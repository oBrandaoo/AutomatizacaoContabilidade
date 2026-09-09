from io import BytesIO
import json
from pathlib import Path
import sqlite3
import zipfile

from openpyxl import load_workbook
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
import pytest

from app import create_app
from nfse import COLUMNS, MONEY_FIELDS, classify, money, parse_pdf, parse_xml

ROOT = Path(__file__).resolve().parents[1]
HEADERS = {'X-NFSe-App': 'local'}


def pdf_bytes(lines=()):
    writer = PdfWriter()
    page = writer.add_blank_page(width=595, height=842)
    if lines:
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        stream = DecodedStreamObject()
        escaped = [line.replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)') for line in lines]
        stream.set_data(('BT /F1 12 Tf 40 780 Td 18 TL '+ ' T* '.join(f'({line}) Tj' for line in escaped)+' ET').encode('latin-1'))
        page[NameObject('/Contents')] = writer._add_object(stream)
    result = BytesIO()
    writer.write(result)
    return result.getvalue()


@pytest.fixture
def app(tmp_path):
    return create_app({'TESTING': True, 'DATABASE': str(tmp_path / 'test.sqlite3')})


@pytest.fixture
def client(app):
    client = app.test_client()
    response = client.post('/api/clients', json={'name': 'Cliente de demonstração', 'document': '12.345.678/0001-95', 'city': 'Santa Rita do Sapucaí / MG'}, headers=HEADERS)
    assert response.status_code == 201
    return client


def upload(client, data, name='nota.xml', client_id=1):
    return client.post('/api/imports', data={'client_id': str(client_id), 'files': (BytesIO(data), name)}, headers=HEADERS)


def sample(name='nota-nacional.xml'):
    return (ROOT / 'examples' / name).read_bytes()


def test_national_and_abrasf_data_and_missing_values():
    national = parse_xml(sample())[0]
    assert national['number'] == '1001'
    assert national['provider_name'] == 'Cliente de demonstração'
    assert national['service_value'] == '2500.00'
    assert national['iss'] == '75.00'
    assert national['pis'] is None
    assert national['status'] == 'Não verificada'
    assert classify(national, '12345678000195') == 'Emitida'
    abrasf = parse_xml(sample('nota-abrasf.xml'))[0]
    assert abrasf['net_value'] == '1193.10'
    assert abrasf['pis'] == '0.00'
    assert abrasf['inss'] is None
    assert abrasf['iss_withheld'] == 'Sim'
    assert classify(abrasf, '12345678000195') == 'Recebida'


def test_reimport_deduplicates_and_original_download_is_exact(client):
    assert upload(client, sample()).json['added'] == 1
    assert upload(client, sample(), 'renomeada.xml').json['duplicates'] == 1
    listing = client.get('/api/invoices').json
    assert listing['stats']['total'] == 1
    detail = client.get('/api/invoices/1').json
    attachment = client.get(f"/api/attachments/{detail['attachments'][0]['id']}")
    assert attachment.data == sample()
    assert 'attachment;' in attachment.headers['Content-Disposition']


def test_client_isolation_filters_and_no_duplicate_across_clients(client):
    client.post('/api/clients', json={'name': 'Segundo cliente', 'document': '98765432000198', 'city': 'Itajubá / MG'}, headers=HEADERS)
    upload(client, sample())
    upload(client, sample(), client_id=2)
    assert client.get('/api/invoices?client_id=1&direction=Emitida').json['stats']['total'] == 1
    assert client.get('/api/invoices?client_id=2&direction=Recebida').json['stats']['total'] == 1
    assert client.get('/api/invoices?start=2026-10-01').json['stats']['total'] == 0
    assert client.get('/api/invoices?date_field=competence&end=2026-08-31').json['stats']['total'] == 0
    assert client.get('/api/invoices?start=2026-10-01&end=2026-09-01').status_code == 400
    assert client.get('/api/invoices?q=inexistente').json['stats']['total'] == 0


def test_zip_partial_errors_and_no_path_extraction(client, tmp_path):
    stream = BytesIO()
    with zipfile.ZipFile(stream, 'w') as archive:
        archive.writestr('../../escape.xml', sample())
        archive.writestr('recebida.xml', sample('nota-abrasf.xml'))
        archive.writestr('broken.xml', '<NFSe>')
        archive.writestr('arquivo.exe', b'not executable')
    result = upload(client, stream.getvalue(), 'lote.zip').json
    assert result['added'] == 2
    assert result['errors'] == 2
    assert client.get('/api/imports').json[0]['errors'] == 2
    assert not (tmp_path / 'escape.xml').exists()


def test_zip_bomb_limit_and_dtd_are_rejected(client):
    stream = BytesIO()
    with zipfile.ZipFile(stream, 'w', zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('big.xml', b'x' * (21 * 1024 * 1024))
    assert upload(client, stream.getvalue(), 'big.zip').json['errors'] == 1
    malicious = b'<!DOCTYPE foo [<!ENTITY xxe SYSTEM "file:///etc/passwd">]><NFSe><infNFSe>&xxe;</infNFSe></NFSe>'
    assert upload(client, malicious).json['errors'] == 1
    assert client.get('/api/invoices').json['stats']['total'] == 0


def test_dps_and_product_invoice_are_not_accepted(client):
    for data in (b'<DPS><infDPS><nDPS>123</nDPS></infDPS></DPS>', b'<nfeProc><NFe><infNFe/></NFe></nfeProc>'):
        assert upload(client, data).json['errors'] == 1


def test_pdf_text_extraction_and_scanned_manual_review(client):
    text = pdf_bytes(['Numero da NFS-e: 321', 'Data de emissao: 05/09/2026', 'Valor dos servicos: R$ 1.234,56'])
    parsed = parse_pdf(text)[0]
    assert parsed['number'] == '321'
    assert parsed['service_value'] == '1234.56'
    assert parsed['review_status'] == 'Pendente'
    assert upload(client, pdf_bytes(), 'digitalizada.pdf').json['added'] == 1
    detail = client.get('/api/invoices/1').json
    assert 'OCR não incluído' in detail['notes']
    assert detail['service_value'] is None
    invalid = client.patch('/api/invoices/1', json={'review_status': 'Conferida'}, headers=HEADERS)
    assert invalid.status_code == 400
    changed = client.patch('/api/invoices/1', json={'number': '55', 'issued_at': '2026-09-03', 'service_value': '1.000,15', 'provider_document': '12345678000195', 'recipient_name': 'Tomador', 'review_status': 'Conferida'}, headers=HEADERS)
    assert changed.status_code == 200
    assert client.get('/api/invoices/1').json['service_value'] == '1000.15'


@pytest.mark.parametrize('pdf_first', [True, False])
def test_pdf_xml_pairing_preserves_xml_as_primary(client, pdf_first):
    key = parse_xml(sample())[0]['access_key']
    assert len(key) == 50
    pdf = pdf_bytes([key, 'Numero da NFS-e: 999', 'Valor dos servicos: R$ 99,00'])
    files = [(pdf, 'nota.pdf'), (sample(), 'nota.xml')]
    if not pdf_first:
        files.reverse()
    for content, name in files:
        assert upload(client, content, name).status_code == 200
    items = client.get('/api/invoices').json['items']
    assert len(items) == 1
    assert items[0]['number'] == '1001'
    assert items[0]['service_value'] == '2500.00'
    detail = client.get(f"/api/invoices/{items[0]['id']}").json
    assert len(detail['attachments']) == 2


def test_multi_invoice_xml_atomic_when_one_invoice_invalid(client):
    good = sample('nota-abrasf.xml').decode('utf-8').split('?>',1)[1]
    bad = good.replace('<Numero>2001</Numero>', '<Numero></Numero>')
    result = upload(client, f'<Lista>{good}{bad}</Lista>'.encode()).json
    assert result['errors'] == 1
    assert client.get('/api/invoices').json['stats']['total'] == 0


def test_export_order_headers_literals_dates_money_and_selection(client):
    upload(client, sample())
    upload(client, sample('nota-abrasf.xml'))
    client.patch('/api/invoices/1', json={'description': '=HYPERLINK("https://example.com")'}, headers=HEADERS)
    columns = [{'key': 'provider_document', 'label': 'Documento'}, {'key': 'description', 'label': '=Cabeçalho literal'}, {'key': 'service_value', 'label': 'Serviços'}, {'key': 'issued_at', 'label': 'Data'}, {'key': 'pis', 'label': 'PIS'}]
    response = client.post('/api/export', json={'columns': columns, 'filters': {}, 'ids': [1]}, headers=HEADERS)
    assert response.status_code == 200
    wb = load_workbook(BytesIO(response.data))
    ws = wb.worksheets[0]
    assert ws.max_row == 2
    assert [c.value for c in ws[1]] == [c['label'] for c in columns]
    assert ws['B1'].data_type == 's'
    assert ws['A2'].data_type == 's' and ws['A2'].value == '12345678000195'
    assert ws['B2'].data_type == 's' and ws['B2'].value.startswith('=HYPERLINK')
    assert ws['C2'].data_type == 'n' and ws['C2'].value == 2500
    assert ws['D2'].is_date
    assert ws['E2'].value is None
    assert ws.freeze_panes == 'A2'
    assert len(wb.worksheets) == 2


def test_templates_persist_and_validate_columns(app, client):
    columns = [{'key': 'number', 'label': 'Nota'}]
    assert client.post('/api/templates', json={'name': 'Mensal', 'columns': columns}, headers=HEADERS).status_code == 200
    assert create_app({'TESTING': True, 'DATABASE': app.config['DATABASE']}).test_client().get('/api/templates').json[0]['columns'] == columns
    for cols in ([], [{'key':'not_allowed','label':'Inválida'}], columns + columns):
        assert client.post('/api/templates', json={'name':'Inválido','columns':cols}, headers=HEADERS).status_code == 400


def test_cancelled_invoice_excluded_from_totals_but_exportable(client):
    upload(client, sample())
    client.patch('/api/invoices/1', json={'status': 'Cancelada'}, headers=HEADERS)
    stats = client.get('/api/invoices').json['stats']
    assert stats['total'] == 1
    assert stats['amounts']['Emitida'] == '0.00'
    assert client.get('/api/invoices?status=Cancelada').json['stats']['total'] == 1


def test_unmatched_client_cannot_mark_invoice_reviewed(client):
    upload(client, sample().replace(b'12345678000195', b'11111111000111'))
    assert client.get('/api/invoices').json['items'][0]['direction'] == 'A conferir'
    assert client.patch('/api/invoices/1', json={'review_status':'Conferida'}, headers=HEADERS).status_code == 400


def test_local_origin_guard_and_api_limits(client):
    assert client.post('/api/clients', json={}).status_code == 403
    assert client.post('/api/clients', json={}, headers={**HEADERS, 'Origin': 'https://external.example'}).status_code == 403
    assert client.get('/', headers={'Host':'hostile.example'}).status_code == 400
    assert client.get('/').status_code == 200
    assert 'frame-ancestors' in client.get('/').headers['Content-Security-Policy']
    assert upload(client, b'', 'empty.xml').json['errors'] == 1
    assert client.post('/api/imports', data={'client_id':'1'}, headers=HEADERS).status_code == 400


@pytest.mark.parametrize('raw', ['NaN','Infinity','-Infinity','abc','99999999999999999999999'])
def test_invalid_money_rejected(raw):
    with pytest.raises(ValueError):
        money(raw)
