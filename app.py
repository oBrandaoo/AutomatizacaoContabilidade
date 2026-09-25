"""Aplicação local de conferência de NFS-e. Execute com python app.py."""
from datetime import date, datetime
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import re
import sqlite3
import zipfile

from flask import Flask, g, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from werkzeug.exceptions import HTTPException

from nfse import (COLUMNS, DEFAULT_COLUMNS, MONEY_FIELDS, classify, document,
                  identity, iso_date, money, parse_pdf, parse_xml)

ROOT = Path(__file__).resolve().parent
MAX_FILE = 20 * 1024 * 1024
SCHEMA = '''
CREATE TABLE IF NOT EXISTS clients (
 id INTEGER PRIMARY KEY, name TEXT NOT NULL, document TEXT NOT NULL UNIQUE, city TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS invoices (
 id INTEGER PRIMARY KEY, client_id INTEGER NOT NULL REFERENCES clients(id),
 identity TEXT NOT NULL, data TEXT NOT NULL, created_at TEXT NOT NULL,
 UNIQUE(client_id, identity));
CREATE TABLE IF NOT EXISTS blobs (hash TEXT PRIMARY KEY, content BLOB NOT NULL);
CREATE TABLE IF NOT EXISTS attachments (
 id INTEGER PRIMARY KEY, invoice_id INTEGER NOT NULL REFERENCES invoices(id),
 hash TEXT NOT NULL REFERENCES blobs(hash), name TEXT NOT NULL, kind TEXT NOT NULL,
 UNIQUE(invoice_id, hash));
CREATE TABLE IF NOT EXISTS imports (
 id INTEGER PRIMARY KEY, client_id INTEGER NOT NULL REFERENCES clients(id),
 created_at TEXT NOT NULL, result TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS templates (
 id INTEGER PRIMARY KEY, name TEXT NOT NULL UNIQUE, columns TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS issuer_settings (
 id INTEGER PRIMARY KEY CHECK (id = 1), name TEXT NOT NULL,
 document TEXT NOT NULL, city TEXT NOT NULL, city_code TEXT NOT NULL,
 municipal_registration TEXT NOT NULL, tax_regime TEXT NOT NULL,
 updated_at TEXT NOT NULL);
'''


def create_app(test_config=None):
    app = Flask(__name__)
    app.config.update(DATABASE=os.environ.get('NFSE_DATABASE', str(ROOT / 'data' / 'notas.sqlite3')), MAX_CONTENT_LENGTH=60 * 1024 * 1024,
                      TRUSTED_HOSTS=['localhost', '127.0.0.1', '[::1]'])
    if test_config:
        app.config.update(test_config)
    Path(app.config['DATABASE']).parent.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(app.config['DATABASE']) as conn:
        conn.executescript(SCHEMA)
        conn.execute('PRAGMA journal_mode=WAL')

    def db():
        if 'db' not in g:
            g.db = sqlite3.connect(app.config['DATABASE'], timeout=30)
            g.db.row_factory = sqlite3.Row
            g.db.execute('PRAGMA foreign_keys=ON')
        return g.db

    @app.teardown_appcontext
    def close_db(_error):
        connection = g.pop('db', None)
        if connection is not None:
            connection.close()

    @app.before_request
    def local_request_guard():
        if request.method in ('POST', 'PUT', 'PATCH', 'DELETE'):
            if request.headers.get('X-NFSe-App') != 'local':
                return jsonify(error='Requisição não autorizada. Recarregue a aplicação.'), 403
            origin = request.headers.get('Origin')
            if origin and origin != request.host_url.rstrip('/'):
                return jsonify(error='Origem da requisição não permitida.'), 403

    @app.after_request
    def headers(response):
        response.headers['X-Content-Type-Options'] = 'nosniff'
        response.headers['X-Frame-Options'] = 'DENY'
        response.headers['Referrer-Policy'] = 'no-referrer'
        response.headers['Content-Security-Policy'] = "default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' data:; object-src 'none'; frame-ancestors 'none'; base-uri 'self'; form-action 'self'"
        if request.path.startswith('/api/'):
            response.headers['Cache-Control'] = 'no-store'
        return response

    @app.errorhandler(ValueError)
    def bad_value(exc):
        return jsonify(error=str(exc)), 400

    @app.errorhandler(HTTPException)
    def http_error(exc):
        message = 'Envio muito grande. Limite de 60 MB por lote.' if exc.code == 413 else exc.description
        return jsonify(error=message), exc.code

    def body():
        result = request.get_json()
        if not isinstance(result, dict):
            raise ValueError('Formato de solicitação inválido.')
        return result

    def client_by_id(client_id):
        row = db().execute('SELECT * FROM clients WHERE id=?', (client_id,)).fetchone()
        if row is None:
            raise ValueError('Selecione um cliente cadastrado.')
        return dict(row)

    def invoice_by_id(invoice_id):
        row = db().execute('SELECT i.*, c.name AS client_name, c.document AS client_document FROM invoices i JOIN clients c ON c.id=i.client_id WHERE i.id=?', (invoice_id,)).fetchone()
        if row is None:
            raise ValueError('Nota não encontrada.')
        result = json.loads(row['data'])
        result.update(id=row['id'], client_id=row['client_id'], client_name=row['client_name'], client_document=row['client_document'])
        result['attachments'] = [dict(r) for r in db().execute('SELECT id,name,kind FROM attachments WHERE invoice_id=? ORDER BY id DESC', (invoice_id,))]
        return result

    def filtered_invoices(filters):
        rows = db().execute('SELECT i.*, c.name AS client_name FROM invoices i JOIN clients c ON c.id=i.client_id ORDER BY i.id DESC').fetchall()
        result = []
        date_field = filters.get('date_field', 'issued_at')
        if date_field not in ('issued_at', 'competence'):
            raise ValueError('Campo de data inválido.')
        start, end = filters.get('start', ''), filters.get('end', '')
        if start:
            start = iso_date(start)
        if end:
            end = iso_date(end)
        if start and end and start > end:
            raise ValueError('A data inicial deve ser anterior à data final.')
        query = str(filters.get('q', '')).casefold().strip()
        for row in rows:
            item = json.loads(row['data'])
            item.update(id=row['id'], client_id=row['client_id'], client_name=row['client_name'])
            if filters.get('client_id') and str(row['client_id']) != str(filters['client_id']):
                continue
            if any(filters.get(key) and item.get(key) != filters[key] for key in ('direction', 'review_status', 'status')):
                continue
            day = item.get(date_field, '')
            if (start and (not day or day < start)) or (end and (not day or day > end)):
                continue
            if query and query not in ' '.join(str(item.get(k) or '') for k in ('number', 'client_name', 'provider_name', 'recipient_name', 'provider_document', 'recipient_document', 'description', 'access_key')).casefold():
                continue
            item.pop('raw_text', None)
            result.append(item)
        return result

    @app.get('/')
    def index():
        return render_template('index.html')

    @app.get('/api/config')
    def config():
        return jsonify(columns=COLUMNS, default_columns=DEFAULT_COLUMNS, version='1.0.0', integration='pending')

    @app.get('/api/clients')
    def clients():
        return jsonify([dict(r) for r in db().execute('SELECT c.*, COUNT(i.id) AS invoice_count FROM clients c LEFT JOIN invoices i ON i.client_id=c.id GROUP BY c.id ORDER BY c.name')])

    @app.get('/api/issuer')
    def issuer():
        row = db().execute('SELECT * FROM issuer_settings WHERE id=1').fetchone()
        return jsonify(dict(row) if row else None)

    @app.put('/api/issuer')
    def save_issuer():
        data = body()
        name = str(data.get('name', '')).strip()
        doc = document(data.get('document', ''))
        city = str(data.get('city', '')).strip()
        city_code = str(data.get('city_code', '')).strip()
        registration = str(data.get('municipal_registration', '')).strip()
        tax_regime = str(data.get('tax_regime', '')).strip()
        if not name or len(name) > 150:
            raise ValueError('Informe a razão social do emitente (até 150 caracteres).')
        if len(doc) != 14 or not doc[-2:].isdigit():
            raise ValueError('Informe o CNPJ do emitente com 14 posições.')
        if not city or len(city) > 150:
            raise ValueError('Informe a cidade e UF do estabelecimento emitente.')
        if not re.fullmatch(r'\d{7}', city_code):
            raise ValueError('Informe o código IBGE do município com 7 dígitos.')
        if len(registration) > 30:
            raise ValueError('A inscrição municipal deve ter até 30 caracteres.')
        if tax_regime not in ('MEI', 'Simples Nacional', 'Lucro Presumido/Real'):
            raise ValueError('Selecione o regime tributário do emitente.')
        with db():
            db().execute('''INSERT INTO issuer_settings
                (id,name,document,city,city_code,municipal_registration,tax_regime,updated_at)
                VALUES(1,?,?,?,?,?,?,?) ON CONFLICT(id) DO UPDATE SET
                name=excluded.name, document=excluded.document, city=excluded.city,
                city_code=excluded.city_code, municipal_registration=excluded.municipal_registration,
                tax_regime=excluded.tax_regime, updated_at=excluded.updated_at''',
                (name, doc, city, city_code, registration, tax_regime, datetime.now().isoformat(timespec='seconds')))
        return jsonify(ok=True)

    @app.post('/api/clients')
    def add_client():
        data = body()
        name = str(data.get('name', '')).strip()
        doc = document(data.get('document', ''))
        city = str(data.get('city', '')).strip()
        if not name or len(name) > 150:
            raise ValueError('Informe o nome do cliente (até 150 caracteres).')
        if len(doc) != 14 or not doc[-2:].isdigit():
            raise ValueError('Informe um CNPJ com 14 posições, incluindo os dois dígitos finais.')
        if not city or len(city) > 150:
            raise ValueError('Informe a cidade e UF.')
        try:
            with db():
                cursor = db().execute('INSERT INTO clients(name,document,city) VALUES(?,?,?)', (name, doc, city))
        except sqlite3.IntegrityError as exc:
            raise ValueError('Já existe um cliente com esse CNPJ.') from exc
        return jsonify(id=cursor.lastrowid), 201

    @app.get('/api/invoices')
    def invoices():
        items = filtered_invoices(request.args)
        sums = {}
        for direction in ('Emitida', 'Recebida'):
            sums[direction] = str(sum((Decimal(i['service_value']) for i in items if i['direction'] == direction and i['service_value'] is not None and i['status'] != 'Cancelada'), Decimal('0.00')))
        stats = dict(total=len(items), pending=sum(i['review_status'] == 'Pendente' for i in items), amounts=sums)
        page = max(1, int(request.args.get('page', 1)))
        return jsonify(items=items[(page-1)*50:page*50], stats=stats, page=page, pages=max(1, (len(items)+49)//50))

    @app.get('/api/invoices/<int:invoice_id>')
    def detail(invoice_id):
        return jsonify(invoice_by_id(invoice_id))

    @app.patch('/api/invoices/<int:invoice_id>')
    def edit_invoice(invoice_id):
        original = invoice_by_id(invoice_id)
        data = body()
        allowed = set(COLUMNS) - {'client_name', 'source', 'access_key', 'direction'}
        for key in allowed:
            if key in data:
                if key in MONEY_FIELDS:
                    original[key] = money(data[key])
                elif key in ('issued_at', 'competence'):
                    original[key] = iso_date(data[key])
                else:
                    original[key] = str(data[key] or '').strip()[:10000]
        for key in ('provider_document', 'recipient_document'):
            original[key] = document(original[key])
        original['direction'] = classify(original, original['client_document'])
        if original['status'] not in ('Não verificada', 'Normal', 'Cancelada', 'Substituída'):
            raise ValueError('Situação inválida.')
        if original['review_status'] not in ('Pendente', 'Conferida'):
            raise ValueError('Conferência inválida.')
        if original['iss_withheld'] not in ('', 'Sim', 'Não'):
            raise ValueError('Retenção de ISS inválida.')
        if original['review_status'] == 'Conferida':
            if not original['number'] or not original['issued_at'] or original['service_value'] is None or original['direction'] == 'A conferir':
                raise ValueError('Para conferir, preencha número, emissão, valor e o CPF/CNPJ do prestador/tomador que vincula a nota ao cliente.')
        for key in ('id', 'client_id', 'client_name', 'client_document', 'attachments'):
            original.pop(key, None)
        with db():
            db().execute('UPDATE invoices SET data=? WHERE id=?', (json.dumps(original, ensure_ascii=False), invoice_id))
        return jsonify(ok=True)

    def import_document(client, filename, content):
        filename = filename.replace('\\', '/').split('/')[-1][:200]
        ext = Path(filename).suffix.lower()
        if ext not in ('.xml', '.pdf'):
            raise ValueError('Formato não suportado. Use XML, PDF ou ZIP com XML/PDF.')
        if not content or len(content) > MAX_FILE:
            raise ValueError('Arquivo vazio ou maior que 20 MB.')
        digest = sha256(content).hexdigest()
        existing = db().execute('SELECT 1 FROM attachments a JOIN invoices i ON i.id=a.invoice_id WHERE i.client_id=? AND a.hash=?', (client['id'], digest)).fetchone()
        if existing:
            return dict(file=filename, state='duplicate', message='Arquivo já importado.', added=0, linked=0)
        parsed = parse_xml(content) if ext == '.xml' else parse_pdf(content)
        added = linked = 0
        with db():
            db().execute('INSERT OR IGNORE INTO blobs(hash,content) VALUES(?,?)', (digest, content))
            for item in parsed:
                item['direction'] = classify(item, client['document'])
                key = identity(item) or 'file:' + digest
                row = db().execute('SELECT * FROM invoices WHERE client_id=? AND identity=?', (client['id'], key)).fetchone()
                if row:
                    invoice_id = row['id']
                    previous = json.loads(row['data'])
                    if ext == '.xml' and previous['source'] == 'PDF':
                        item['raw_text'] = previous.get('raw_text', '')
                        item['notes'] = 'XML incorporado ao PDF pela chave de acesso. Confira os dados estruturados.'
                        db().execute('UPDATE invoices SET data=? WHERE id=?', (json.dumps(item, ensure_ascii=False), invoice_id))
                    elif ext == '.xml':
                        previous['review_status'] = 'Pendente'
                        previous['notes'] += '\nOutra versão de XML anexada. Compare os originais antes de conferir.'
                        if item['status'] == 'Cancelada':
                            previous['status'] = 'Cancelada'
                        db().execute('UPDATE invoices SET data=? WHERE id=?', (json.dumps(previous, ensure_ascii=False), invoice_id))
                    elif not previous.get('raw_text'):
                        previous['raw_text'] = item['raw_text']
                        db().execute('UPDATE invoices SET data=? WHERE id=?', (json.dumps(previous, ensure_ascii=False), invoice_id))
                    linked += 1
                else:
                    cursor = db().execute('INSERT INTO invoices(client_id,identity,data,created_at) VALUES(?,?,?,?)',
                                          (client['id'], key, json.dumps(item, ensure_ascii=False), datetime.now().isoformat(timespec='seconds')))
                    invoice_id = cursor.lastrowid
                    added += 1
                db().execute('INSERT OR IGNORE INTO attachments(invoice_id,hash,name,kind) VALUES(?,?,?,?)', (invoice_id, digest, filename, ext[1:].upper()))
        return dict(file=filename, state='success', message=f'{added} nota(s) adicionada(s), {linked} arquivo(s) vinculado(s).', added=added, linked=linked)

    @app.post('/api/imports')
    def upload():
        client = client_by_id(request.form.get('client_id'))
        files = request.files.getlist('files')
        if not files or len(files) > 150:
            raise ValueError('Selecione entre 1 e 150 arquivos.')
        results = []
        processed = expanded = 0

        def process(name, data):
            nonlocal processed, expanded
            processed += 1
            expanded += len(data)
            if processed > 150 or expanded > 100 * 1024 * 1024:
                raise ValueError('Lote excede 150 arquivos ou 100 MB descompactados.')
            try:
                results.append(import_document(client, name, data))
            except Exception as exc:
                app.logger.info('Importação recusada: %s', type(exc).__name__)
                message = str(exc) if isinstance(exc, ValueError) else 'Arquivo inválido, danificado ou com estrutura não permitida.'
                results.append(dict(file=name[:200], state='error', message=message, added=0, linked=0))

        for file in files:
            name = file.filename or 'arquivo'
            try:
                data = file.read(MAX_FILE + 1)
                if len(data) > MAX_FILE:
                    raise ValueError('Arquivo maior que 20 MB.')
                if Path(name).suffix.lower() == '.zip':
                    with zipfile.ZipFile(BytesIO(data)) as archive:
                        members = [m for m in archive.infolist() if not m.is_dir()]
                        if len(members) + processed > 150 or sum(m.file_size for m in members) + expanded > 100 * 1024 * 1024:
                            raise ValueError('ZIP excede 150 arquivos ou 100 MB descompactados.')
                        if any(m.file_size > MAX_FILE or m.flag_bits & 1 for m in members):
                            raise ValueError('ZIP contém arquivo maior que 20 MB ou protegido por senha.')
                        if not members:
                            raise ValueError('ZIP vazio.')
                        for member in members:
                            # Read into memory; never extract archive paths onto disk.
                            process(member.filename, archive.read(member))
                else:
                    process(name, data)
            except Exception as exc:
                message = str(exc) if isinstance(exc, ValueError) else 'Não foi possível ler o ZIP. Confira se o arquivo está íntegro.'
                results.append(dict(file=name[:200], state='error', message=message, added=0, linked=0))
        summary = dict(results=results, added=sum(r['added'] for r in results), linked=sum(r['linked'] for r in results),
                       errors=sum(r['state'] == 'error' for r in results), duplicates=sum(r['state'] == 'duplicate' for r in results))
        with db():
            db().execute('INSERT INTO imports(client_id,created_at,result) VALUES(?,?,?)',
                         (client['id'], datetime.now().isoformat(timespec='seconds'), json.dumps(summary, ensure_ascii=False)))
        return jsonify(summary)

    @app.get('/api/imports')
    def history():
        return jsonify([dict(id=r['id'], client_name=r['name'], created_at=r['created_at'], **json.loads(r['result']))
                        for r in db().execute('SELECT imports.*, clients.name FROM imports JOIN clients ON clients.id=imports.client_id ORDER BY imports.id DESC LIMIT 100')])

    @app.get('/api/attachments/<int:attachment_id>')
    def download(attachment_id):
        row = db().execute('SELECT a.*, b.content FROM attachments a JOIN blobs b ON b.hash=a.hash WHERE a.id=?', (attachment_id,)).fetchone()
        if row is None:
            raise ValueError('Arquivo não encontrado.')
        return send_file(BytesIO(row['content']), as_attachment=True, download_name=row['name'], mimetype='application/pdf' if row['kind'] == 'PDF' else 'application/xml')

    def validate_columns(columns):
        if not isinstance(columns, list) or not 1 <= len(columns) <= len(COLUMNS):
            raise ValueError('Selecione pelo menos uma coluna.')
        seen = set()
        result = []
        for col in columns:
            if not isinstance(col, dict) or col.get('key') not in COLUMNS or col['key'] in seen:
                raise ValueError('Coluna inválida ou repetida.')
            label = str(col.get('label', '')).strip()
            if not label or len(label) > 100:
                raise ValueError('Dê um nome de até 100 caracteres para cada coluna.')
            result.append(dict(key=col['key'], label=label))
            seen.add(col['key'])
        return result

    @app.get('/api/templates')
    def templates():
        return jsonify([dict(id=r['id'], name=r['name'], columns=json.loads(r['columns'])) for r in db().execute('SELECT * FROM templates ORDER BY name')])

    @app.post('/api/templates')
    def save_template():
        data = body()
        name = str(data.get('name', '')).strip()
        if not name or len(name) > 80:
            raise ValueError('Informe um nome de até 80 caracteres para o modelo.')
        columns = validate_columns(data.get('columns'))
        with db():
            db().execute('INSERT INTO templates(name,columns) VALUES(?,?) ON CONFLICT(name) DO UPDATE SET columns=excluded.columns', (name, json.dumps(columns, ensure_ascii=False)))
        return jsonify(ok=True)

    @app.delete('/api/templates/<int:template_id>')
    def delete_template(template_id):
        with db():
            db().execute('DELETE FROM templates WHERE id=?', (template_id,))
        return jsonify(ok=True)

    @app.post('/api/export')
    def export():
        data = body()
        columns = validate_columns(data.get('columns'))
        filters = data.get('filters', {})
        if not isinstance(filters, dict):
            raise ValueError('Filtros inválidos.')
        items = filtered_invoices(filters)
        selected = data.get('ids')
        if selected is not None:
            if not isinstance(selected, list) or any(type(i) is not int for i in selected):
                raise ValueError('Seleção de notas inválida.')
            ids = set(selected)
            items = [i for i in items if i['id'] in ids]
        if not items:
            raise ValueError('Não há notas para exportar com essa seleção.')
        if len(items) > 100000:
            raise ValueError('Limite de 100 mil notas por exportação. Restrinja o período.')
        wb = Workbook()
        ws = wb.active
        ws.title = 'Notas de serviço'
        ws.append([c['label'] for c in columns])
        for cell in ws[1]:
            cell.data_type = 's'
            cell.fill = PatternFill('solid', fgColor='174C40')
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(vertical='center')
        ws.row_dimensions[1].height = 28
        for item in items:
            row_values = []
            for col in columns:
                val = item.get(col['key'])
                if val is not None and col['key'] in MONEY_FIELDS:
                    val = Decimal(val)
                elif val and col['key'] in ('issued_at', 'competence'):
                    val = date.fromisoformat(val)
                row_values.append(val)
            ws.append(row_values)
            for index, col in enumerate(columns, 1):
                cell = ws.cell(ws.max_row, index)
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # Identifiers and formula-like user text stay literal.
                if col['key'] in MONEY_FIELDS:
                    cell.number_format = '#,##0.00'
                elif col['key'] in ('issued_at', 'competence'):
                    cell.number_format = 'dd/mm/yyyy'
                cell.alignment = Alignment(vertical='top', wrap_text=col['key'] in ('description', 'notes'))
        for index, col in enumerate(columns, 1):
            ws.column_dimensions[ws.cell(1, index).column_letter].width = 48 if col['key'] in ('description', 'notes', 'access_key') else 24
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        info = wb.create_sheet('Sobre a exportação')
        info.append(['Gerada em', datetime.now().isoformat(timespec='seconds')])
        info.append(['Quantidade de notas', len(items)])
        info.append(['Notas pendentes de conferência', sum(i['review_status'] == 'Pendente' for i in items)])
        info.append(['Origem dos dados', 'Arquivos importados e revisões locais. Não houve consulta fiscal online.'])
        info.append(['Valores ausentes', 'Célula vazia significa dado não informado; não equivale a zero.'])
        info.column_dimensions['A'].width = 36
        info.column_dimensions['B'].width = 90
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return send_file(buffer, as_attachment=True, download_name=f'notas-servico-{date.today().isoformat()}.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    return app


if __name__ == '__main__':
    from waitress import serve
    print('Notas | Abra http://127.0.0.1:8000 — Ctrl+C para encerrar.', flush=True)
    serve(create_app(), host='127.0.0.1', port=8000, threads=4, max_request_body_size=60 * 1024 * 1024)
